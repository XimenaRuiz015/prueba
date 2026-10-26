import os
import sys
import re
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from pypdf import PdfReader
from docx import Document
import chromadb
from chromadb.utils import embedding_functions

from app.database.db import SessionLocal, init_db
from app.database.models import Proyecto
from app.database.crud import crear_proyecto

DOCUMENTS_DIR = BASE_DIR / "documents"
CHROMA_DIR = BASE_DIR / "data" / "chroma_db"

def extraer_texto_pdf(ruta_archivo: Path) -> list[dict]:
    paginas = []
    try:
        reader = PdfReader(str(ruta_archivo))
        for num_pag, pagina in enumerate(reader.pages, start=1):
            texto = pagina.extract_text()
            if texto and texto.strip():
                paginas.append({"texto": texto.strip(), "pagina": num_pag})
    except Exception as e:
        print(f"   [!] Error leyendo PDF {ruta_archivo.name}: {e}", flush=True)
    return paginas

def extraer_texto_docx(ruta_archivo: Path) -> list[dict]:
    secciones = []
    try:
        doc = Document(str(ruta_archivo))
        texto_completo = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables[:20]:
            for row in table.rows[:30]:
                fila = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if fila:
                    texto_completo.append(" | ".join(fila))
        if texto_completo:
            secciones.append({"texto": "\n\n".join(texto_completo), "pagina": 1})
    except Exception as e:
        print(f"   [!] Error leyendo Word {ruta_archivo.name}: {e}", flush=True)
    return secciones

def dividir_en_chunks(texto: str, tamano_chunk: int = 1000, solapamiento: int = 150) -> list[str]:
    if len(texto) <= tamano_chunk:
        return [texto]
    chunks = []
    inicio = 0
    while inicio < len(texto):
        chunks.append(texto[inicio:inicio + tamano_chunk])
        inicio += (tamano_chunk - solapamiento)
    return chunks

def procesar_e_ingestar():
    init_db()
    db = SessionLocal()
    client = chromadb.PersistentClient(path=str(CHROMA_DIR))
    emb_fn = embedding_functions.SentenceTransformerEmbeddingFunction(
        model_name="paraphrase-multilingual-MiniLM-L12-v2"
    )
    coleccion = client.get_or_create_collection(
        name="equigrup_documentacion",
        embedding_function=emb_fn,
        metadata={"hnsw:space": "cosine"}
    )

    print("=" * 60, flush=True)
    print(">>> INICIANDO INGESTA DE CONOCIMIENTO PARA EQUIGRUP", flush=True)
    print("=" * 60, flush=True)

    carpetas = {
        "informes_tecnicos": DOCUMENTS_DIR / "informes_tecnicos",
        "procedimientos": DOCUMENTS_DIR / "procedimientos"
    }

    total_documentos = 0
    total_chunks_nuevos = 0
    proyectos_creados = 0

    for tipo_doc, ruta_carpeta in carpetas.items():
        if not ruta_carpeta.exists():
            continue
        archivos = [
            f for f in ruta_carpeta.iterdir()
            if f.is_file() and not f.name.startswith("~") and not f.name.startswith(".") and f.suffix.lower() in [".pdf", ".docx", ".doc"]
        ]
        print(f"\n[+] Procesando carpeta: [{tipo_doc}] ({len(archivos)} archivos encontrados)", flush=True)
        
        for idx_arch, archivo in enumerate(archivos, start=1):
            nombre_archivo = archivo.name
            extension = archivo.suffix.lower()
            print(f"   [{idx_arch}/{len(archivos)}] Ingeriendo: {nombre_archivo[:55]}...", flush=True)
            
            if extension == ".pdf":
                paginas = extraer_texto_pdf(archivo)
            elif extension in [".docx", ".doc"]:
                paginas = extraer_texto_docx(archivo)
            else:
                continue

            if not paginas:
                continue

            total_documentos += 1
            texto_primer_parrafo = ""
            ids_doc, docs_doc, metas_doc = [], [], []

            for item in paginas:
                num_pag = item["pagina"]
                chunks = dividir_en_chunks(item["texto"])
                if not texto_primer_parrafo and item["texto"]:
                    texto_primer_parrafo = item["texto"][:400].replace("\n", " ")
                for c_idx, ch in enumerate(chunks):
                    clean_name = re.sub(r"[^a-zA-Z0-9]", "_", nombre_archivo)
                    chunk_id = f"{tipo_doc}_{clean_name}_p{num_pag}_c{c_idx}"
                    ids_doc.append(chunk_id)
                    docs_doc.append(ch)
                    metas_doc.append({"documento": nombre_archivo, "tipo": tipo_doc, "pagina": num_pag})

            if ids_doc:
                coleccion.upsert(ids=ids_doc, documents=docs_doc, metadatas=metas_doc)
                total_chunks_nuevos += len(ids_doc)

            if tipo_doc == "informes_tecnicos":
                titulo_proyecto = nombre_archivo.rsplit(".", 1)[0].replace("_", " ").title()
                existe = db.query(Proyecto).filter(Proyecto.titulo == titulo_proyecto).first()
                if not existe:
                    crear_proyecto(db, {
                        "titulo": titulo_proyecto,
                        "cliente": "Cliente Registrado en Informe",
                        "ubicacion": "Planta / Campo",
                        "categoria": "Informe Técnico",
                        "estado": "Completado",
                        "maquinaria": "Ver detalles en informe",
                        "descripcion": texto_primer_parrafo or f"Informe técnico correspondiente al archivo {nombre_archivo}.",
                        "creado_por": "Ingesta Automática"
                    })
                    proyectos_creados += 1

    db.close()
    print("\n" + "=" * 60, flush=True)
    print("[OK] INGESTA MASIVA COMPLETADA CON EXITO", flush=True)
    print(f"   • Total de documentos procesados: {total_documentos}", flush=True)
    print(f"   • Total de fragmentos agregados al RAG: {total_chunks_nuevos}", flush=True)
    print(f"   • Nuevos proyectos registrados en la Base de Datos: {proyectos_creados}", flush=True)
    print("=" * 60, flush=True)

if __name__ == "__main__":
    procesar_e_ingestar()
